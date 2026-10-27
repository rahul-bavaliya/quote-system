"""Synthetic exports in exactly the same layout as the real BI exports (junk rows above the
header, Power-BI column names, fixed-width part descriptions, parts split over yearly sheets).

Used by the test-suite and by `make sample` so the system can be tried without company data.

    python -m quote_system.sample_data data/raw-sample
"""
from __future__ import annotations

import random
import sys
from datetime import date, timedelta
from pathlib import Path

WO_HEADERS = ["Branch[Branch Country]", "ServiceOrder[Wo Number]", "ServiceOrder[Wo Segment]",
              "ServiceOrder[ServiceOrderID]", "ServiceOrder[Contract Type]", "ServiceOrder[ShopOffsite Description]",
              "ServiceOrder[Make]", "ServiceOrder[Model]", "ServiceOrder[Meter]", "ServiceOrder[Description]",
              "ServiceOrder Sales Type[Description]", "ServiceOrder Status[Description]", "ServiceOrder[Open Date]",
              "ServiceOrder[Invoice Date]", "ServiceOrder[EQ Labor Hours]", "Branch[Branch Name]",
              "[M_Actual_ServiceOrderLaborHrsTotal]"]
PART_HEADERS = ["ServiceOrder Product Detail[Product Number]", "ServiceOrder Product Detail[Description]",
                "ServiceOrder Product Detail[ServiceOrderPartsShipQty]", "ServiceOrder Product Detail[ServiceOrderID]"]

MODELS = [("JD", "333G", "Compact Construction Equipment", 1.0), ("JD", "333P", "Compact Construction Equipment", 1.0),
          ("JD", "544P", "C&F", 1.2), ("JD", "S780", "AG", 1.4), ("JD", "872GP", "C&F", 1.3)]
# (description choices, parts [(pn, desc, qty)], hours median)
JOBS = {
    "battery": (["REPLACE BATTERY", "REPLACE BATTERIES", "R&R BATTERY", "REMPLACER BATTERIE"],
                [("TY25879B", "WET CHAR", 1), ("TY26783B", "WET CHAR", 0)], 1.5),
    "starter": (["REPLACE STARTER", "STARTER R&R", "STARTER NOT ENGAGING - REPLACE"],
                [("SE501406", "STARTER", 1), ("CRSE501406", "LESS COR", 1), ("19M7184", "CAP SCRE", 2)], 3.0),
    "ecu": (["PROGRAM NEW ECU", "REPLACE ECU"], [("SE502800", "ENGINE C", 1)], 2.5),
    "pm500": (["500 HOUR SERVICE", "500 HR SERVICE", "PERFORM 500HR PM"],
              [("AT365869", "OIL FILT", 1), ("RE541922", "FUEL FIL", 1), ("TY26674", "PLUS-50", 4), ("WDF", "WASTE DI", 1)], 4.0),
    "hose": (["REPLACE HYDRAULIC HOSE", "BLOWN HOSE ON BOOM"], [("AT312345", "HOSE", 1), ("51M7044", "O-RING", 2)], 2.5),
    "travel": (["TRAVEL", "SERVICE CALL", "TRAVEL TO SITE"], [], 2.0),
}


def _desc(pn: str, d: str, code: str = "") -> str:
    return f"{pn:<16}{d:<8}{code}".rstrip()


def generate(out_dir: Path, n_wo: int = 2500, seed: int = 7) -> tuple[Path, Path]:
    from openpyxl import Workbook

    rnd = random.Random(seed)
    out_dir.mkdir(parents=True, exist_ok=True)
    wo_rows, part_rows = [], {}
    start = date(2024, 1, 1)
    for wo in range(1, n_wo + 1):
        make, model, div, mult = rnd.choice(MODELS)
        nseg = rnd.choice([1, 1, 1, 2])
        for seg in range(1, nseg + 1):
            kind = rnd.choices(list(JOBS) + ["nostart_batt", "nostart_starter", "batt_starter"],
                               weights=[30, 12, 6, 20, 12, 15, 8, 4, 3])[0]
            parts = []
            if kind == "nostart_batt":
                desc, parts, med = "MACHINE NOT STARTING", JOBS["battery"][1], 2.0
            elif kind == "nostart_starter":
                desc, parts, med = "NO START", JOBS["starter"][1], 3.5
            elif kind == "batt_starter":
                desc, parts, med = "REPLACE BATTERY AND STARTER", JOBS["battery"][1] + JOBS["starter"][1], 4.0
            else:
                choices, parts, med = JOBS[kind]
                desc = rnd.choice(choices)
            act = round(max(0.25, rnd.lognormvariate(0, 0.35) * med * mult) * 4) / 4
            opened = start + timedelta(days=rnd.randint(0, 900))
            sid = f"S1{wo:08d}{seg:03d}0101"
            eq = act if rnd.random() < 0.4 else 0
            wo_rows.append(["CA", wo, seg, sid, div, rnd.choice(["Shop", "Off site"]), make, model,
                            rnd.randint(100, 9000), desc, rnd.choice(["Charge", "Cash", "Warranty", "Internal"]),
                            "Invoiced", opened, opened + timedelta(days=rnd.randint(1, 20)), eq,
                            rnd.choice(["Regina", "Saskatoon", "Edmonton CCE"]), act])
            year = str(opened.year)
            for pn, d, q in parts:
                if q == 0 and rnd.random() < 0.7:
                    continue
                part_rows.setdefault(year, []).append([pn, _desc(pn, d, rnd.choice(["", "", "CY"])), max(q, 1), sid])

    wo_path, parts_path = out_dir / "workorders.xlsx", out_dir / "parts.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = "BTLServiceOrder"
    ws.append(["Data returned for BTLServiceOrder."])
    ws.append([])
    ws.append(WO_HEADERS)
    for r in wo_rows:
        ws.append(r)
    wb.save(wo_path)
    wb = Workbook()
    wb.remove(wb.active)
    for year in sorted(part_rows):
        ws = wb.create_sheet(year)
        ws.append([None, None, None, None, "Data returned for BTLServiceOrder."])
        ws.append([])
        ws.append(PART_HEADERS)
        for r in part_rows[year]:
            ws.append(r)
    wb.save(parts_path)
    return wo_path, parts_path


if __name__ == "__main__":
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "data/raw-sample")
    w, p = generate(target)
    print(f"wrote {w} and {p}")
